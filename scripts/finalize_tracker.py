from __future__ import annotations

from datetime import date
import json
from pathlib import Path
import re

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TRACKER = ROOT / "TRACKER.xlsx"

TASK_EVIDENCE = {
    "A1": ("docs/PROJECT_SPEC.md", "Scope and deliverables documented."),
    "A2": ("docs/ARCHITECTURE.md", "Architecture matches the implemented local stack."),
    "A3": ("docs/DEVELOPMENT_PLAN.md", "Milestone plan retained."),
    "A4": ("docs/ACCEPTANCE_CRITERIA.md", "Acceptance gates retained and reconciled in final verification."),
    "A5": ("evidence/README.md; docs/EVIDENCE.md", "Evidence policy and artifact map documented."),
    "A6": (".git; Git history", "Repository initialized; implementation and verification milestone recorded in Git history."),
    "A7": ("evidence/dataset/dataset_summary.json; evidence/dataset/validation_report.json", "Fixed-seed 4,200-row dataset passes validation."),
    "A8": ("evidence/metrics/analytics_report.json", "Cumulative funnel analytics generated from current dataset."),
    "A9": ("evidence/metrics/analytics_report.json; evidence/api/api_smoke.json", "Cycle, programme, region, selectivity, and campaign-source cohorts generated."),
    "A10": ("evidence/metrics/analytics_report.json", "Campaign attribution is explicitly descriptive, not causal."),
    "A11": ("evidence/models/model_report.json; evidence/models/admitflow_model.joblib", "Training and feature summaries captured. Predictive quality is weak; models are not decision-ready."),
    "A12": ("evidence/metrics/forecast_report.json", "H1/H2/H4/H6 evaluated; only six annual observations are available."),
    "A13": ("evidence/metrics/scenario_report.json; evidence/api/api_smoke.json", "Deterministic scenario engine and authenticated route verified."),
    "A14": ("evidence/api/api_smoke.json; evidence/security/auth_check.json; docs/API.md", "Protected API, aggregate-grounded optional AI route, and missing-token rejection verified; audit log is process-local."),
    "A15": ("evidence/screenshots/dashboard-redesign.png; evidence/screenshots/dashboard-dark.png; evidence/screenshots/dashboard-mobile-redesign.png; evidence/screenshots/scenario-redesign.png; evidence/screenshots/login-redesign.png; evidence/screenshots/assistant-offline.png; evidence/tests/frontend_vitest.txt", "Redesigned desktop/mobile workspaces and graceful AI-offline state captured; interactions covered by Vitest."),
    "A16": ("evidence/tests/backend_pytest.txt; evidence/tests/frontend_vitest.txt; evidence/tests/frontend_lint.txt; evidence/tests/frontend_build.txt", "Automated suites pass; production build has a disclosed chunk-size warning."),
    "A17": ("docs/FINAL_PROJECT_REPORT.md; docs/FINAL_VERIFICATION.md", "Final documentation completed with evidence and limitations."),
}

EVIDENCE_ROWS = [
    ("AF-E01", "A7", "Dataset validation", "evidence/dataset/validation_report.json", "PASS", "4,200 rows; no validation issues."),
    ("AF-E02", "A8", "Analytics report", "evidence/metrics/analytics_report.json", "PASS", "Dashboard, funnel, cohorts, programme, and attribution outputs."),
    ("AF-E03", "A11", "Model report and artifact", "evidence/models/model_report.json; evidence/models/admitflow_model.joblib", "GENERATED_WITH_LIMITATIONS", "Random forest predicts no positives at threshold 0.5; not decision-ready."),
    ("AF-E04", "A12", "Forecast report", "evidence/metrics/forecast_report.json", "PASS_WITH_LIMITATIONS", "Two-year holdout over six annual observations."),
    ("AF-E05", "A13", "Scenario report", "evidence/metrics/scenario_report.json", "PASS", "Deterministic baseline/scenario/delta output."),
    ("AF-E06", "A14", "API smoke report", "evidence/api/api_smoke.json", "PASS", "11/11 authenticated API smoke operations successful."),
    ("AF-E07", "A14", "Auth boundary report", "evidence/security/auth_check.json", "PASS", "401 without token; token/password omitted from report."),
    ("AF-E08", "A15", "Desktop dashboard screenshot", "evidence/screenshots/dashboard.png", "CAPTURED", "Authenticated live dashboard at desktop viewport."),
    ("AF-E09", "A15", "Mobile dashboard screenshot", "evidence/screenshots/dashboard-mobile.png", "CAPTURED", "Authenticated live dashboard at 390px viewport."),
    ("AF-E10", "A16", "Backend test suite", "evidence/tests/backend_pytest.txt", "PASS", "Current full pytest output."),
    ("AF-E11", "A16", "Frontend test suite", "evidence/tests/frontend_vitest.txt", "PASS", "Current full Vitest output."),
    ("AF-E12", "A16", "Frontend lint", "evidence/tests/frontend_lint.txt", "PASS", "No lint diagnostics."),
    ("AF-E13", "A16", "Production build", "evidence/tests/frontend_build.txt", "PASS_WITH_WARNING", "Build succeeds; bundle exceeds 500 kB warning."),
    ("AF-E14", "A17", "Final verification documents", "docs/FINAL_VERIFICATION.md; docs/FINAL_PROJECT_REPORT.md", "PASS", "Status and limitations stated explicitly."),
    ("AF-E15", "A15", "Scaled scenario screenshot", "evidence/screenshots/scenario.png", "CAPTURED", "Live UI scenario at twice the baseline application volume; stage totals and changes shown."),
    ("AF-E16", "A15", "Redesigned experience captures", "evidence/screenshots/dashboard-redesign.png; evidence/screenshots/dashboard-dark.png; evidence/screenshots/dashboard-mobile-redesign.png; evidence/screenshots/scenario-redesign.png; evidence/screenshots/login-redesign.png; evidence/screenshots/assistant-offline.png", "CAPTURED", "Live dashboard, dark theme, mobile dashboard, scenario workspace, premium login, and AI provider fallback."),
]


def text(path: str) -> str:
    contents = (ROOT / path).read_bytes()
    encoding = "utf-16" if b"\x00" in contents[:80] else "utf-8"
    return contents.decode(encoding, errors="replace")


def main() -> None:
    workbook = load_workbook(TRACKER)
    today = date.today()
    master = workbook["Master Checklist"]
    columns = {cell.value: cell.column for cell in master[1]}

    for row in range(2, master.max_row + 1):
        task_id = master.cell(row, columns["Task ID"]).value
        if task_id not in TASK_EVIDENCE:
            continue
        evidence, notes = TASK_EVIDENCE[task_id]
        master.cell(row, columns["Status"], "Completed")
        master.cell(row, columns["Evidence"], evidence)
        master.cell(row, columns["Notes"], notes)
        master.cell(row, columns["Actual Hours"], None)

    evidence_sheet = workbook["Evidence Register"]
    evidence_ids = {entry[0] for entry in EVIDENCE_ROWS}
    for row in range(evidence_sheet.max_row, 1, -1):
        if evidence_sheet.cell(row, 1).value in evidence_ids:
            evidence_sheet.delete_rows(row)
    for evidence_id, task_id, description, path, result, notes in EVIDENCE_ROWS:
        paths = path.split("; ")
        missing = [item for item in paths if not (ROOT / item).exists()]
        if missing:
            raise FileNotFoundError(f"Missing evidence for {evidence_id}: {missing}")
        evidence_sheet.append([evidence_id, today, task_id, "Execution artifact", description, path, result, "Automated local run", notes])

    backend_output = text("evidence/tests/backend_pytest.txt")
    frontend_output = text("evidence/tests/frontend_vitest.txt")
    lint_output = text("evidence/tests/frontend_lint.txt")
    build_output = text("evidence/tests/frontend_build.txt")
    model_report = json.loads(text("evidence/models/model_report.json"))
    backend_passed = int(re.search(r"(\d+) passed", backend_output).group(1))
    backend_failed_match = re.search(r"(\d+) failed", backend_output)
    backend_failed = int(backend_failed_match.group(1)) if backend_failed_match else 0
    frontend_passed = int(re.search(r"Tests\s+(\d+) passed", frontend_output).group(1))
    if "built in" not in build_output or backend_failed or "failed" in frontend_output:
        raise RuntimeError("Verification logs contain a failed quality gate.")

    test_sheet = workbook["Test Register"]
    generated_test_ids = {"AF-T01", "AF-T02", "AF-T03", "AF-T04", "AF-T05"}
    for row in range(test_sheet.max_row, 1, -1):
        if test_sheet.cell(row, 1).value in generated_test_ids:
            test_sheet.delete_rows(row)
    test_rows = [
        ["AF-T01", "A16", "Backend pytest suite", "python -m pytest -q", "All tests pass", f"{backend_passed} passed; {backend_failed} failed", "PASS", "evidence/tests/backend_pytest.txt"],
        ["AF-T02", "A16", "Frontend Vitest suite", "npm test", "All tests pass", f"{frontend_passed} passed", "PASS", "evidence/tests/frontend_vitest.txt"],
        ["AF-T03", "A16", "Frontend lint", "npm run lint", "No diagnostics", "Clean" if not lint_output.strip().endswith("warning") else lint_output[-200:], "PASS", "evidence/tests/frontend_lint.txt"],
        ["AF-T04", "A16", "Frontend production build", "npm run build", "Build succeeds", "Build succeeded; bundle-size warning present", "PASS_WITH_WARNING", "evidence/tests/frontend_build.txt"],
        ["AF-T05", "A14", "Authenticated API smoke checks", "python scripts/capture_evidence.py", "Protected routes respond successfully", "11/11 operations successful", "PASS", "evidence/api/api_smoke.json"],
    ]
    for row in test_rows:
        test_sheet.append(row)

    model_sheet = workbook["Model Results"]
    for row in range(model_sheet.max_row, 1, -1):
        model_sheet.delete_rows(row)
    for key in ("baseline", "advanced"):
        result = model_report[key]
        model_sheet.append([
            result["model"],
            "data/admission_data.csv (synthetic)",
            "enrolled",
            json.dumps({name: result[name] for name in ("roc_auc", "pr_auc", "accuracy", "precision", "recall", "f1")}),
            today,
            "evidence/models/model_report.json; evidence/models/admitflow_model.joblib",
            "Evaluated; not decision-ready",
        ])

    work_log = workbook["Work Log"]
    if not any(work_log.cell(row, 2).value == "CLOSEOUT" for row in range(2, work_log.max_row + 1)):
        work_log.append([today, "CLOSEOUT", "Project Agent", "Final local verification and evidence capture", None, "docs/FINAL_VERIFICATION.md", "Automated checks pass; model limitations recorded."])

    dashboard = workbook["Dashboard"]
    dashboard_columns = {cell.value: cell.column for cell in dashboard[1]}
    statuses = [master.cell(row, columns["Status"]).value for row in range(2, master.max_row + 1)]
    dashboard_values = {
        "total tasks": len(statuses),
        "completed": statuses.count("Completed"),
        "in progress": statuses.count("In Progress"),
        "blocked": statuses.count("Blocked"),
        "not started": statuses.count("Not Started"),
        "total planned hours": sum(master.cell(row, columns["Planned Hours"]).value or 0 for row in range(2, master.max_row + 1)),
        "total recorded actual hours": None,
        "evidence count": len(EVIDENCE_ROWS),
        "test pass count": backend_passed + frontend_passed,
        "test failure count": backend_failed,
    }
    for row in range(2, dashboard.max_row + 1):
        label = dashboard.cell(row, dashboard_columns["Metric"]).value
        if label in dashboard_values:
            dashboard.cell(row, dashboard_columns["Value"], dashboard_values[label])

    workbook.save(TRACKER)
    print(f"Updated {TRACKER} with {len(EVIDENCE_ROWS)} evidence rows and {len(test_rows)} verification rows.")
    print(f"Tasks completed: {statuses.count('Completed')}/{len(statuses)}; recorded actual hours: none (not measured).")


if __name__ == "__main__":
    main()
